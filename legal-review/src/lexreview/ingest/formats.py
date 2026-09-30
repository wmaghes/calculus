"""Format detection and parsing. Runs INSIDE the sandboxed worker process.

Input: raw bytes. Output: a ParseResult with page-mapped text. Every failure
maps to a coverage status plus a reason code. Nothing here logs, and
exceptions are converted to reason codes before leaving the worker.

Page mapping by format:
    PDF     native pages (1-based). Pages with no text layer are OCR'd.
    Image   one page per frame, OCR'd.
    DOCX    DOCX has no fixed pages. Units are split at Word's own
            last-rendered page breaks when present (locator says
            "Word layout, approx."), else at explicit page breaks, else in
            blocks of 40 paragraphs. The locator always gives paragraph numbers.
    XLSX    one unit per 200 rows of each sheet; locator "Sheet!rows a-b".
    Email   unit 1 = headers + body. Attachments are returned as child files
            and ingested as their own documents with parent_id set.
    Text    split on form feeds, else blocks of 60 lines.
"""

from __future__ import annotations

import email
import email.policy
import io
import re
import zipfile
from dataclasses import dataclass, field
from html.parser import HTMLParser

DOCX_PARA_BLOCK = 40
XLSX_ROW_BLOCK = 200
TEXT_LINE_BLOCK = 60
MAX_CELLS = 2_000_000
MAX_CHILDREN = 200


@dataclass
class PageOut:
    page_no: int
    locator: str
    text: str
    ocr: bool = False
    ocr_conf: float | None = None


@dataclass
class ParseResult:
    kind: str
    status: str = "indexed"
    reason: str | None = None
    pages: list[PageOut] = field(default_factory=list)
    meta: dict = field(default_factory=dict)
    children: list[tuple[str, bytes]] = field(default_factory=list)


class Reject(Exception):
    def __init__(self, status: str, reason: str):
        super().__init__(reason)
        self.status, self.reason = status, reason


# ---------------------------------------------------------------- detection
_EMAIL_HDR = re.compile(rb"^(received|from|to|subject|date|message-id|mime-version|return-path|delivered-to):", re.I | re.M)
_IMAGE_MAGIC = (b"\x89PNG\r\n\x1a\n", b"\xff\xd8\xff", b"II*\x00", b"MM\x00*", b"GIF87a", b"GIF89a", b"BM")


def detect(data: bytes) -> str:
    head = data[:2048]
    if b"%PDF-" in data[:1024]:
        return "pdf"
    if head.startswith(b"PK\x03\x04"):
        return "zip"
    if head.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
        return "ole"
    if any(head.startswith(m) for m in _IMAGE_MAGIC):
        return "image"
    if b"\x00" in head:
        return "binary"
    first_lines = head.split(b"\n\n", 1)[0].split(b"\r\n\r\n", 1)[0]
    if len(_EMAIL_HDR.findall(first_lines)) >= 2:
        return "email"
    return "text"


# ---------------------------------------------------------------- OCR
def ocr_image(img, timeout: int) -> tuple[str, float]:
    import pytesseract

    try:
        d = pytesseract.image_to_data(img, output_type=pytesseract.Output.DICT, timeout=timeout)
    except RuntimeError as exc:  # pytesseract raises RuntimeError on timeout
        raise Reject("timeout", "ocr_timeout") from exc
    except pytesseract.TesseractError as exc:
        raise Reject("ocr_failed", "tesseract_error") from exc
    lines: dict[tuple, list[str]] = {}
    confs = []
    for i, word in enumerate(d["text"]):
        if not word.strip():
            continue
        key = (d["block_num"][i], d["par_num"][i], d["line_num"][i])
        lines.setdefault(key, []).append(word)
        c = float(d["conf"][i])
        if c >= 0:
            confs.append(c)
    text = "\n".join(" ".join(ws) for _, ws in sorted(lines.items()))
    return text, (sum(confs) / len(confs) if confs else 0.0)


# ---------------------------------------------------------------- PDF
def parse_pdf(data: bytes, limits: dict) -> ParseResult:
    import pypdfium2 as pdfium

    res = ParseResult(kind="pdf")
    try:
        pdf = pdfium.PdfDocument(data)
    except pdfium.PdfiumError as exc:
        if getattr(exc, "err_code", None) == 4:
            raise Reject("password_protected", "pdf_password") from exc
        raise Reject("corrupted", "pdf_unreadable") from exc
    n = len(pdf)
    if n == 0:
        raise Reject("corrupted", "pdf_no_pages")
    if n > limits["max_pages"]:
        raise Reject("too_large", "page_limit")
    for i in range(n):
        try:
            page = pdf[i]
            text = page.get_textpage().get_text_range()
        except pdfium.PdfiumError:
            text = ""
            page = None
        text = text.replace("\r\n", "\n").replace("\r", "\n")
        if len(text.strip()) >= limits["ocr_min_chars"] or page is None:
            res.pages.append(PageOut(i + 1, f"p. {i + 1}", text))
            continue
        img = page.render(scale=300 / 72).to_pil()
        ocr_text, conf = ocr_image(img, limits["ocr_timeout"])
        res.pages.append(PageOut(i + 1, f"p. {i + 1}", ocr_text, ocr=True, ocr_conf=round(conf, 1)))
    return res


# ---------------------------------------------------------------- images
def parse_image(data: bytes, limits: dict) -> ParseResult:
    from PIL import Image, ImageSequence

    Image.MAX_IMAGE_PIXELS = limits["max_image_pixels"]
    res = ParseResult(kind="image")
    try:
        img = Image.open(io.BytesIO(data))  # lazy: reads the header only
        # Pillow only *raises* above 2x MAX_IMAGE_PIXELS (it merely warns
        # between 1x and 2x), so enforce the limit ourselves before decoding.
        if img.size[0] * img.size[1] > limits["max_image_pixels"]:
            raise Reject("rejected_unsafe", "image_decompression_bomb")
        img.load()
    except Reject:
        raise
    except Image.DecompressionBombError as exc:
        raise Reject("rejected_unsafe", "image_decompression_bomb") from exc
    except Exception as exc:  # noqa: BLE001 - Pillow raises many types
        raise Reject("corrupted", "image_unreadable") from exc
    for i, frame in enumerate(ImageSequence.Iterator(img)):
        if i >= limits["max_pages"]:
            raise Reject("too_large", "page_limit")
        text, conf = ocr_image(frame.convert("RGB"), limits["ocr_timeout"])
        res.pages.append(PageOut(i + 1, f"image frame {i + 1}", text, ocr=True, ocr_conf=round(conf, 1)))
    return res


# ---------------------------------------------------------------- ZIP-based
def check_zip(data: bytes, limits: dict) -> zipfile.ZipFile:
    try:
        zf = zipfile.ZipFile(io.BytesIO(data))
        infos = zf.infolist()
    except zipfile.BadZipFile as exc:
        raise Reject("corrupted", "zip_unreadable") from exc
    total = sum(i.file_size for i in infos)
    packed = max(1, sum(i.compress_size for i in infos))
    if total > limits["max_zip_uncompressed"] or total / packed > limits["max_zip_ratio"]:
        raise Reject("rejected_unsafe", "zip_bomb_suspected")
    return zf


def parse_zip(data: bytes, limits: dict) -> ParseResult:
    zf = check_zip(data, limits)
    names = set(zf.namelist())
    if "word/document.xml" in names:
        return parse_docx(data)
    if "xl/workbook.xml" in names:
        return parse_xlsx(data)
    raise Reject("unsupported", "zip_archive_not_supported")


_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _docx_blocks(doc):
    """Yield (text, rendered_break_before, explicit_break_before) per paragraph
    or table row, in body order."""
    for el in doc.element.body.iterchildren():
        if el.tag == _W + "p":
            rendered = el.find(f".//{_W}lastRenderedPageBreak") is not None
            explicit = any(br.get(_W + "type") == "page" for br in el.iter(_W + "br"))
            text = "".join(t.text or "" for t in el.iter(_W + "t"))
            yield text, rendered, explicit
        elif el.tag == _W + "tbl":
            for tr in el.iter(_W + "tr"):
                cells = ["".join(t.text or "" for t in tc.iter(_W + "t")) for tc in tr.iter(_W + "tc")]
                yield " | ".join(cells), False, False


def parse_docx(data: bytes) -> ParseResult:
    import docx

    try:
        d = docx.Document(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        raise Reject("corrupted", "docx_unreadable") from exc
    blocks = list(_docx_blocks(d))
    has_rendered = any(b[1] for b in blocks)
    has_explicit = any(b[2] for b in blocks)
    res = ParseResult(kind="docx")
    units: list[tuple[int, int, list[str]]] = []  # (first_para, last_para, lines)
    cur: list[str] = []
    first = 1
    for idx, (text, rendered, explicit) in enumerate(blocks, start=1):
        brk = (has_rendered and rendered) or (not has_rendered and has_explicit and explicit)
        if (brk and cur) or (not has_rendered and not has_explicit and len(cur) >= DOCX_PARA_BLOCK):
            units.append((first, idx - 1, cur))
            cur, first = [], idx
        cur.append(text)
    if cur or not units:
        units.append((first, max(first, len(blocks)), cur))
    for n, (a, b, lines) in enumerate(units, start=1):
        if has_rendered:
            loc = f"p. {n} (Word layout, approx.), ¶{a}–¶{b}"
        elif has_explicit:
            loc = f"section {n} (explicit page breaks), ¶{a}–¶{b}"
        else:
            loc = f"¶{a}–¶{b}"
        res.pages.append(PageOut(n, loc, "\n".join(lines)))
    res.meta["page_basis"] = "word_rendered" if has_rendered else ("explicit_breaks" if has_explicit else "paragraph_blocks")
    return res


def parse_xlsx(data: bytes) -> ParseResult:
    import openpyxl

    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise Reject("corrupted", "xlsx_unreadable") from exc
    res = ParseResult(kind="xlsx")
    cells = 0
    n = 0
    for ws in wb.worksheets:
        buf: list[str] = []
        start = 1
        for r, row in enumerate(ws.iter_rows(values_only=True), start=1):
            cells += len(row)
            if cells > MAX_CELLS:
                raise Reject("too_large", "cell_limit")
            buf.append("\t".join("" if v is None else str(v) for v in row))
            if len(buf) >= XLSX_ROW_BLOCK:
                n += 1
                res.pages.append(PageOut(n, f"{ws.title}!rows {start}–{r}", "\n".join(buf)))
                buf, start = [], r + 1
        if buf:
            n += 1
            res.pages.append(PageOut(n, f"{ws.title}!rows {start}–{start + len(buf) - 1}", "\n".join(buf)))
    res.meta["note"] = "formula_cells_show_cached_values"
    return res


def parse_ole(data: bytes) -> ParseResult:
    # Encrypted OOXML (password-protected .docx/.xlsx) is an OLE container
    # holding "EncryptionInfo" and "EncryptedPackage" streams.
    if "EncryptionInfo".encode("utf-16-le") in data or "EncryptedPackage".encode("utf-16-le") in data:
        raise Reject("password_protected", "office_encrypted")
    raise Reject("unsupported", "legacy_office_or_msg")


# ---------------------------------------------------------------- email
class _HTMLText(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.out: list[str] = []
        self._skip = 0

    def handle_starttag(self, tag, attrs):
        if tag in ("script", "style"):
            self._skip += 1
        if tag in ("br", "p", "div", "tr", "li"):
            self.out.append("\n")

    def handle_endtag(self, tag):
        if tag in ("script", "style") and self._skip:
            self._skip -= 1

    def handle_data(self, data):
        if not self._skip:
            self.out.append(data)


def html_to_text(html: str) -> str:
    p = _HTMLText()
    p.feed(html)
    return re.sub(r"\n{3,}", "\n\n", "".join(p.out)).strip()


def parse_email(data: bytes) -> ParseResult:
    try:
        msg = email.message_from_bytes(data, policy=email.policy.default)
    except Exception as exc:  # noqa: BLE001
        raise Reject("corrupted", "email_unreadable") from exc
    res = ParseResult(kind="email")
    hdr = []
    for h in ("From", "To", "Cc", "Date", "Subject"):
        v = msg.get(h)
        if v is not None:
            hdr.append(f"{h}: {v}")
            res.meta[h.lower()] = str(v)
    body_part = msg.get_body(preferencelist=("plain", "html"))
    body = ""
    if body_part is not None:
        try:
            content = body_part.get_content()
        except (LookupError, UnicodeError):
            content = body_part.get_payload(decode=True).decode("utf-8", "replace")
        body = html_to_text(content) if body_part.get_content_type() == "text/html" else content
    text = "\n".join(hdr) + "\n\n" + body.replace("\r\n", "\n")
    res.pages.append(PageOut(1, "email headers and body", text))
    for att in msg.iter_attachments():
        if len(res.children) >= MAX_CHILDREN:
            res.meta["attachments_truncated"] = True
            break
        payload = att.get_payload(decode=True) or b""
        name = att.get_filename() or "attachment"
        res.children.append((re.sub(r"[^A-Za-z0-9._ \-]", "_", name)[:120], payload))
    res.meta["attachments"] = len(res.children)
    return res


# ---------------------------------------------------------------- text
def parse_text(data: bytes) -> ParseResult:
    try:
        text = data.decode("utf-8")
    except UnicodeDecodeError:
        text = data.decode("latin-1")
        printable = sum(ch.isprintable() or ch in "\n\t\r" for ch in text[:10000])
        if printable < 0.95 * min(len(text), 10000):
            raise Reject("unsupported", "binary_or_unknown_encoding")
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    res = ParseResult(kind="text")
    if "\f" in text:
        for i, part in enumerate(text.split("\f"), start=1):
            res.pages.append(PageOut(i, f"p. {i} (form feed)", part))
        return res
    lines = text.split("\n")
    for i in range(0, max(1, len(lines)), TEXT_LINE_BLOCK):
        block = lines[i:i + TEXT_LINE_BLOCK]
        res.pages.append(PageOut(len(res.pages) + 1, f"lines {i + 1}–{i + len(block)}", "\n".join(block)))
    return res


# ---------------------------------------------------------------- entry
def parse(data: bytes, limits: dict) -> ParseResult:
    kind = detect(data)
    try:
        if kind == "pdf":
            res = parse_pdf(data, limits)
        elif kind == "zip":
            res = parse_zip(data, limits)
        elif kind == "ole":
            res = parse_ole(data)
        elif kind == "image":
            res = parse_image(data, limits)
        elif kind == "email":
            res = parse_email(data)
        elif kind == "text":
            res = parse_text(data)
        else:
            raise Reject("unsupported", "unknown_binary_format")
    except Reject as r:
        return ParseResult(kind=kind, status=r.status, reason=r.reason)
    except MemoryError:
        return ParseResult(kind=kind, status="rejected_unsafe", reason="memory_limit")
    except Exception as exc:  # noqa: BLE001 - map anything else to a code
        return ParseResult(kind=kind, status="parse_error", reason="unexpected_" + type(exc).__name__)
    # OCR outcome -> document status.
    ocr_pages = [p for p in res.pages if p.ocr]
    if ocr_pages:
        failed = [p for p in ocr_pages if not p.text.strip()]
        low = [p for p in ocr_pages if p.text.strip() and (p.ocr_conf or 0) < limits["ocr_low_conf"]]
        if len(failed) == len(res.pages):
            res.status, res.reason = "ocr_failed", "no_text_recovered"
        elif failed or low:
            res.status = "indexed_low_confidence"
            res.reason = f"ocr_failed_pages_{len(failed)}_low_conf_pages_{len(low)}"
    if res.status == "indexed" and not any(p.text.strip() for p in res.pages) and not res.children:
        res.status, res.reason = "indexed_low_confidence", "no_text_content"
    return res
